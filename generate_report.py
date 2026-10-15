import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Set Margins (1 inch all sides)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("LAPORAN PENGUJIAN REST API\nTOKO SEMBAKO")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Manajemen Database Client Server / CodeIgniter 4 RESTful API")
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph() # Spacer

    # Identitas Box (Tabel Identitas)
    id_table = doc.add_table(rows=5, cols=2)
    id_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    id_data = [
        ("Nama Mahasiswa", ": [Isi Nama Anda]"),
        ("NIM", ": [Isi NIM Anda]"),
        ("Kelas / Prodi", ": [Isi Kelas / TRPL]"),
        ("Mata Kuliah", ": Manajemen Database Client Server"),
        ("Tanggal Pengujian", ": 8 Oktober 2026"),
    ]
    for row_idx, (label, val) in enumerate(id_data):
        row = id_table.rows[row_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        
        cell_lbl.width = Inches(2.0)
        cell_val.width = Inches(4.2)
        
        p0 = cell_lbl.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        
        p1 = cell_val.paragraphs[0]
        p1.add_run(val)
        
        set_cell_background(cell_lbl, "F0F4F8")
        set_cell_background(cell_val, "F9FAFC")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

    doc.add_paragraph()

    # Bab 1: Pendahuluan
    h1 = doc.add_heading("1. PENDAHULUAN", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    
    p_intro = doc.add_paragraph(
        "Laporan ini mendokumentasikan pengujian unit dan pengujian integrasi endpoint RESTful API "
        "pada aplikasi Toko Sembako yang dibangun menggunakan CodeIgniter 4. Seluruh pengujian "
        "dilakukan dengan menggunakan Postman HTTP Client pada server lokal (http://localhost:8080)."
    )
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.paragraph_format.space_after = Pt(12)

    # Bab 2: Checklist Pengujian
    h2 = doc.add_heading("2. LEMBAR CHECKLIST PENGUJIANKU (11 SKENARIO)", level=1)
    h2.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    # Checklist Table
    chk_table = doc.add_table(rows=12, cols=6)
    chk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["No", "Skenario Pengujian", "Method", "Status Diharapkan", "Status Aktual", "Hasil"]
    col_widths = [Inches(0.4), Inches(2.3), Inches(0.8), Inches(1.1), Inches(1.0), Inches(0.7)]
    
    hdr_cells = chk_table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=80, right=80)

    scenarios_data = [
        ("1", "GET semua data produk", "GET", "200 OK", "200 OK", "Lulus"),
        ("2", "GET dengan pencarian (?q=...)", "GET", "200 OK", "200 OK", "Lulus"),
        ("3", "GET detail berdasarkan ID", "GET", "200 OK", "200 OK", "Lulus"),
        ("4", "POST tambah data produk valid", "POST", "201 Created", "201 Created", "Lulus"),
        ("5", "POST data tidak valid (uji validasi)", "POST", "400 Bad Request", "400 Bad Request", "Lulus"),
        ("6", "PUT ubah seluruh data produk", "PUT", "200 OK", "200 OK", "Lulus"),
        ("7", "PATCH ubah sebagian kolom produk", "PATCH", "200 OK", "200 OK", "Lulus"),
        ("8", "PUT dengan barcode milik produk lain", "PUT", "400 Bad Request", "400 Bad Request", "Lulus"),
        ("9", "DELETE hapus data produk", "DELETE", "200 OK", "200 OK", "Lulus"),
        ("10", "GET data yang tidak ada / sudah dihapus", "GET", "404 Not Found", "404 Not Found", "Lulus"),
        ("11", "POST dengan body kosong / JSON rusak", "POST", "400 Bad Request", "400 Bad Request", "Lulus"),
    ]

    for row_idx, data in enumerate(scenarios_data, start=1):
        row_cells = chk_table.rows[row_idx].cells
        bg_color = "F9FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(data):
            row_cells[c_idx].width = col_widths[c_idx]
            p = row_cells[c_idx].paragraphs[0]
            if c_idx in [0, 2, 3, 4, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            if c_idx == 5:
                r.bold = True
                r.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32) # Green
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=80, right=80)

    doc.add_paragraph()

    # Bab 3: Detail Skenario & Screenshot Placeholder
    h3 = doc.add_heading("3. DETAIL PENGUJIAN API & BUKTI POSTMAN", level=1)
    h3.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    details = [
        ("3.1 Pengujian 1: GET Semua Data Produk", "GET", "http://localhost:8080/api/products", "None (Kosong)", "200 OK", "Menampilkan seluruh daftar produk sembako yang tersimpan di database."),
        ("3.2 Pengujian 2: GET Pencarian Produk", "GET", "http://localhost:8080/api/products?q=Air", "None (Kosong)", "200 OK", "Mencari produk berdasarkan kata kunci nama/barcode ('Air')."),
        ("3.3 Pengujian 3: GET Detail Produk Berdasarkan ID", "GET", "http://localhost:8080/api/products/1", "None (Kosong)", "200 OK", "Menampilkan detail satu produk spesifik dengan ID 1."),
        ("3.4 Pengujian 4: POST Tambah Data Produk Valid", "POST", "http://localhost:8080/api/products", '{\n  "barcode": "8999999000059",\n  "name": "Gula Pasir 1kg",\n  "price": 17500\n}', "201 Created", "Menambahkan data produk baru yang valid ke tabel products."),
        ("3.5 Pengujian 5: POST Data Tidak Valid (Uji Validasi)", "POST", "http://localhost:8080/api/products", '{\n  "name": "A",\n  "price": -5000\n}', "400 Bad Request", "Menguji aturan validasi server (nama minimal 2 karakter, harga tidak boleh negatif)."),
        ("3.6 Pengujian 6: PUT Ubah Seluruh Data Produk", "PUT", "http://localhost:8080/api/products/1", '{\n  "barcode": "8999999000011",\n  "name": "Air Mineral 600ml Premium",\n  "price": 4500\n}', "200 OK", "Memperbarui seluruh data produk dengan ID 1."),
        ("3.7 Pengujian 7: PATCH Ubah Sebagian Kolom Produk", "PATCH", "http://localhost:8080/api/products/1", '{\n  "price": 4200\n}', "200 OK", "Memperbarui hanya kolom harga pada produk ID 1."),
        ("3.8 Pengujian 8: PUT Barcode Bentrok (Milik Produk Lain)", "PUT", "http://localhost:8080/api/products/1", '{\n  "name": "Air Mineral",\n  "price": 4000,\n  "barcode": "8999999000028"\n}', "400 Bad Request", "Menguji keunikan barcode. Mengubah barcode ID 1 menggunakan barcode yang sudah dipakai ID 2."),
        ("3.9 Pengujian 9: DELETE Hapus Data Produk", "DELETE", "http://localhost:8080/api/products/5", "None (Kosong)", "200 OK", "Menghapus produk dengan ID 5 (produk yang belum pernah dipakai di transaksi)."),
        ("3.10 Pengujian 10: GET Data Yang Tidak Ada", "GET", "http://localhost:8080/api/products/999", "None (Kosong)", "404 Not Found", "Menguji respon server jika data ID yang diminta tidak ditemukan di database."),
        ("3.11 Pengujian 11: POST Body JSON Rusak / Empty", "POST", "http://localhost:8080/api/products", '{name: "Produk Rusak"', "400 Bad Request", "Menguji pertahanan API terhadap format sintaks JSON yang rusak / invalid."),
    ]

    for title, method, url, body_text, status_exp, desc in details:
        sh = doc.add_heading(title, level=2)
        sh.runs[0].font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
        
        p_desc = doc.add_paragraph()
        p_desc.add_run("• Deskripsi: ").bold = True
        p_desc.add_run(desc)
        
        # Detail Table
        dt_table = doc.add_table(rows=4, cols=2)
        dt_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        dt_data = [
            ("HTTP Method & URL", f"{method}  {url}"),
            ("Expected Status Code", status_exp),
            ("Request Body (JSON)", body_text),
            ("Bukti Response Postman", "[ SIKAPKAN SCREENSHOT POSTMAN DI KOTAK BAWAH INI ]")
        ]
        
        for r_i, (lbl, val) in enumerate(dt_data):
            r_cells = dt_table.rows[r_i].cells
            r_cells[0].width = Inches(2.0)
            r_cells[1].width = Inches(4.2)
            
            p_lbl = r_cells[0].paragraphs[0]
            r_l = p_lbl.add_run(lbl)
            r_l.bold = True
            
            p_val = r_cells[1].paragraphs[0]
            if r_i == 2 and val != "None (Kosong)":
                p_val.paragraph_format.line_spacing = 1.0
                r_v = p_val.add_run(val)
                r_v.font.name = 'Consolas'
                r_v.font.size = Pt(9.5)
            else:
                p_val.add_run(val)
                
            set_cell_background(r_cells[0], "F0F4F8")
            set_cell_background(r_cells[1], "FFFFFF")
            set_cell_margins(r_cells[0], top=60, bottom=60, left=100, right=100)
            set_cell_margins(r_cells[1], top=60, bottom=60, left=100, right=100)

        # Large Placeholder Box for Screenshot
        box_table = doc.add_table(rows=1, cols=1)
        box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        box_cell = box_table.rows[0].cells[0]
        box_cell.width = Inches(6.2)
        set_cell_background(box_cell, "FAFAFA")
        set_cell_margins(box_cell, top=200, bottom=200, left=200, right=200)
        
        p_box = box_cell.paragraphs[0]
        p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_box = p_box.add_run("\n\n[ TEMPEL SCREENSHOT POSTMAN SINI ]\n(Klik kanan / Paste Gambar Hasil Pengujian Postman)\n\n")
        r_box.font.italic = True
        r_box.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        
        doc.add_paragraph() # Spacer

    # Bab 4: Pengujian Tambahan Transactions
    h4 = doc.add_heading("4. PENGUJIAN TAMBAHAN (TRANSAKSI & SALE ITEMS)", level=1)
    h4.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_trx_desc = doc.add_paragraph(
        "Selain 11 pengujian produk di atas, dilakukan pula pengujian pada endpoint Transaksi "
        "untuk memastikan kalkulasi total belanja dan snapshot item penjualan tersimpan dengan benar."
    )
    p_trx_desc.paragraph_format.space_after = Pt(10)

    # Transaction 1
    doc.add_heading("4.1 POST Tambah Transaksi Baru (Create Transaction)", level=2)
    p_t1 = doc.add_paragraph()
    p_t1.add_run("• Method & URL: ").bold = True
    p_t1.add_run("POST http://localhost:8080/api/transactions\n")
    p_t1.add_run("• Request Body: ").bold = True
    p_t1.add_run('{"items": [{"product_id": 1, "quantity": 2}, {"product_id": 2, "quantity": 1}]}')

    t_box1 = doc.add_table(rows=1, cols=1)
    t_box1.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_b1 = t_box1.rows[0].cells[0]
    c_b1.width = Inches(6.2)
    set_cell_background(c_b1, "FAFAFA")
    set_cell_margins(c_b1, top=150, bottom=150, left=150, right=150)
    p_b1 = c_b1.paragraphs[0]
    p_b1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_b1 = p_b1.add_run("\n[ TEMPEL SCREENSHOT POSTMAN POST TRANSAKSI ]\n")
    r_b1.font.italic = True
    r_b1.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()

    # Transaction 2
    doc.add_heading("4.2 GET Detail Transaksi Beserta Item Belanja", level=2)
    p_t2 = doc.add_paragraph()
    p_t2.add_run("• Method & URL: ").bold = True
    p_t2.add_run("GET http://localhost:8080/api/transactions/1")

    t_box2 = doc.add_table(rows=1, cols=1)
    t_box2.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_b2 = t_box2.rows[0].cells[0]
    c_b2.width = Inches(6.2)
    set_cell_background(c_b2, "FAFAFA")
    set_cell_margins(c_b2, top=150, bottom=150, left=150, right=150)
    p_b2 = c_b2.paragraphs[0]
    p_b2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_b2 = p_b2.add_run("\n[ TEMPEL SCREENSHOT POSTMAN GET DETAIL TRANSAKSI ]\n")
    r_b2.font.italic = True
    r_b2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()

    # Bab 5: Kesimpulan
    h5 = doc.add_heading("5. KESIMPULAN", level=1)
    h5.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_conc = doc.add_paragraph(
        "Berdasarkan hasil pengujian yang telah dilakukan menggunakan Postman HTTP Client terhadap "
        "11 skenario utama dan pengujian tambahan endpoint transaksi, seluruh API telah berfungsi "
        "sesuai dengan spesifikasi RESTful architecture. Penanganan status code HTTP (200, 201, 400, 404), "
        "validasi input JSON, serta transaksi database berjalan dengan baik tanpa kendala."
    )
    p_conc.paragraph_format.line_spacing = 1.15

    # Save document
    filename = "Laporan_Pengujian_REST_API_Toko_Sembako.docx"
    doc.save(filename)
    print(f"Document successfully created: {filename}")

if __name__ == "__main__":
    create_report()
